# -*- coding: utf-8 -*-
"""Las 21 claves propuestas, con los productos que YA usan esa abreviatura."""
import openpyxl, collections
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
B='/tmp/claude-1001/-home-agentia-odoo-odoo-workspace/40e1928e-dd87-4d8e-b8a3-401e8c22dfcd/scratchpad/'

# ficha -> (clave propuesta, nombre propuesto, grupo)
P = [
 ('2185','COHIE01','Hielera de unicel No. 4','Consumible / embalaje'),
 ('2148','COESP01','Esponja facial de celulosa','Consumible / limpieza'),
 ('2187','COESC01','Escuadra profesional 9 cm sin bisel','Consumible / papelería'),
 ('2176','COPLA01','Puntero láser verde recargable','Consumible / papelería'),
 ('2167','COTIJ01','Tijeras para zurdos 21 cm','Consumible / papelería'),
 ('2164','COCZP01','Cubrezapatos desechables','Consumible / protección'),
 ('2200','COBOT01','Botella boca ancha polietileno ámbar 1000 ml','Consumible de laboratorio'),
 ('2201','COBOT02','Botella boca ancha polietileno ámbar 250 ml','Consumible de laboratorio'),
 ('2175','COJER51','Jeringa de vidrio borosilicato','Consumible de laboratorio'),
 ('2178','COMIP01','Microplaca ELISA 96 pozos','Consumible de laboratorio'),
 ('2153','COPLP04','Placa Petri plástica 60 x 15 mm','Consumible de laboratorio'),
 ('2174','COPRB06','Probeta de plástico 10 ml graduada','Consumible de laboratorio'),
 ('2165','COPUN06','Punta de pipeta 5 ml','Consumible de laboratorio'),
 ('2163','COTUB09','Tubo PCR 0.2 ml tapa plana','Consumible de laboratorio'),
 ('2157','EQBAD02','Báscula de precisión Rhino','Equipo de laboratorio'),
 ('2207','EQCEN01','Centrífuga para laboratorio','Equipo de laboratorio'),
 ('2206','EQTHI01','Higrómetro digital con termómetro y despertador','Equipo de laboratorio'),
 ('2205','EQTHI02','Termohigrómetro digital HTC-1','Equipo de laboratorio'),
 ('2197','EQTRV02','Termómetro Brannan de inmersión parcial','Equipo de laboratorio'),
 ('2155','EQVAG01','Varilla agitadora magnética','Equipo de laboratorio'),
 ('2159','MPREC100','Base caldo Half Fraser granulado 500 g','Materia prima / medio'),
]
# hermanos por abreviatura
herm = collections.defaultdict(list)
pref = None
for ln in open(B+'hermanos.txt'):
    ln = ln.rstrip('\n')
    if ln.startswith('### '):
        pref = ln[4:].strip(); continue
    if '~' in ln and pref:
        f = ln.split('~')
        herm[pref].append((f[0], f[1], f[2]))

MPREC_TOPE = [('MPREC95','Sorbitol 100%'),('MPREC96','Yoduro de Potasio'),
              ('MPREC97','Verde brillante'),('MPREC98','EDTA'),('MPREC99','S17')]

azul=PatternFill('solid',fgColor='1F4E78'); ambar=PatternFill('solid',fgColor='FFF2CC')
verde=PatternFill('solid',fgColor='E8F4E8'); rojo=PatternFill('solid',fgColor='FCE4E4')
borde=Border(*[Side(style='thin',color='BFBFBF')]*4)

def enc(ws,t,sub,cols):
    ws['A1']=t; ws['A1'].font=Font(size=14,bold=True,color='1F4E78')
    ws['A2']=sub; ws['A2'].font=Font(size=9,italic=True,color='808080')
    for i,(h,w) in enumerate(cols,1):
        c=ws.cell(row=4,column=i,value=h); c.font=Font(bold=True,color='FFFFFF'); c.fill=azul
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=borde
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[4].height=32; ws.freeze_panes='A5'

def fila(ws,r,v,rell=None,alto=30):
    for i,x in enumerate(v,1):
        c=ws.cell(row=r,column=i,value=x); c.border=borde
        c.alignment=Alignment(vertical='top',wrap_text=True)
        if rell: c.fill=rell
    ws.row_dimensions[r].height=alto

wb=openpyxl.Workbook()
ws=wb.active; ws.title='21 claves propuestas'
enc(ws,'Las 21 claves propuestas, contra lo que ya existe con esa abreviatura',
    'La columna de la derecha es para que juzgues si la abreviatura encaja. Verde = la abreviatura ya existe y hay con qué comparar. Ámbar = abreviatura NUEVA, no hay precedente.',
    [('Ficha',8),('Producto a codificar',44),('Clave propuesta',15),
     ('¿La abreviatura ya existe?',18),('Productos que YA usan esa abreviatura',72),('Tu visto bueno',16)])
r=5
for fid, clave, nombre, grupo in P:
    pre = clave[:5]
    hs = herm.get(pre, [])
    if pre == 'MPREC':
        texto = 'Hay 99. Los últimos: ' + ' · '.join('%s %s' % (c,n) for c,n in MPREC_TOPE)
        col = rojo
        existe = 'SÍ, y está LLENA'
    elif hs:
        texto = ' · '.join('%s %s' % (c, n) for c, n, _ in hs)
        col = verde; existe = 'sí, %d productos' % len(hs)
    else:
        texto = '(ninguno: sería la primera)'
        col = ambar; existe = 'NO, es nueva'
    fila(ws, r, [fid, nombre, clave, existe, texto, ''], col, 34)
    r += 1

ws2=wb.create_sheet('Ojo con estas')
enc(ws2,'Cuatro que conviene mirar con calma','Las demás son directas.',
    [('Clave propuesta',15),('Producto',40),('Por qué mirarla',96)])
OJO=[('MPREC100','Base caldo Half Fraser granulado 500 g',
      'La serie MPREC llegó a 99 y la normativa dice consecutivo de DOS dígitos. MPREC100 rompe la regla. O se amplía a tres dígitos, o el caldo va en otra abreviatura (es un medio de cultivo, no un reactivo).'),
     ('COJER51','Jeringa de vidrio borosilicato',
      'Las jeringas existentes se numeran por CAPACIDAD, no por consecutivo: COJER03 es de 3 ml, COJER05 de 5 ml, COJER50 de 50 ml. Poner 51 rompe ese patrón. Si la jeringa nueva es de 20 ml, tocaría COJER20.'),
     ('COPRB06','Probeta de plástico 10 ml graduada',
      'Las probetas existentes NO siguen orden de capacidad (25, 100, 500, 250, 1 L), así que 06 va bien. Pero ojo: todas las existentes son de cristalería y esta es de plástico.'),
     ('EQTHI01 y EQTHI02','Higrómetro y Termohigrómetro',
      'Son dos aparatos parecidos con abreviatura nueva. Vale confirmar que de verdad son distintos y no el mismo comprado dos veces.')]
for n,o in enumerate(OJO): fila(ws2,5+n,list(o), ambar, 46)

wb.save(B+'Claves_propuestas_2026-09-29.xlsx')
print('OK: %d claves, %d con abreviatura existente' % (len(P), sum(1 for f in P if herm.get(f[1][:5]))))
