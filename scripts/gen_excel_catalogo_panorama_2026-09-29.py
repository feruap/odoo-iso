# -*- coding: utf-8 -*-
"""Panorama del catalogo interno: los 80 publicados, y que le falta a cada uno."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
B='/tmp/claude-1001/-home-agentia-odoo-odoo-workspace/40e1928e-dd87-4d8e-b8a3-401e8c22dfcd/scratchpad/'

# fichas sin clave que SI entran al almacen -> clave propuesta
ENTRAN = {
 '2148':'COESP01','2153':'COPLP04','2155':'EQVAG01','2157':'EQBAD02','2159':'MPREC100',
 '2163':'COTUB09','2164':'COCZP01','2165':'COPUN06','2167':'COTIJ01','2174':'COPRB06',
 '2175':'COJER51','2176':'COPLA01','2178':'COMIP01','2185':'COHIE01','2187':'COESC01',
 '2197':'EQTRV02','2200':'COBOT01','2201':'COBOT02','2205':'EQTHI02','2206':'EQTHI01',
 '2207':'EQCEN01',
}
# fichas sin clave que NO entran al almacen -> grupo
NO_ENTRAN = {
 '2202':'Mobiliario','2172':'Mobiliario','2203':'Mobiliario','2194':'Mobiliario',
 '2208':'Mobiliario','2198':'Mobiliario','2193':'Mobiliario','2190':'Mobiliario','2184':'Mobiliario',
 '2199':'Cómputo','2196':'Cómputo','2150':'Cómputo','2168':'Cómputo','2156':'Cómputo','2204':'Cómputo',
 '2191':'Instalaciones','2182':'Instalaciones','2322':'Instalaciones','2192':'Instalaciones',
 '2195':'Mantenimiento','2181':'Mantenimiento','2189':'Mantenimiento','2169':'Mantenimiento','2170':'Mantenimiento',
}
filas=[]
for ln in open(B+'cat80.txt'):
    f=ln.rstrip('\n').split('~')
    if len(f)>=8: filas.append(f)   # id, clave, nombre, categoria, flujo, existencia, movs, liga

azul=PatternFill('solid',fgColor='1F4E78'); ambar=PatternFill('solid',fgColor='FFF2CC')
verde=PatternFill('solid',fgColor='E8F4E8'); rojo=PatternFill('solid',fgColor='FCE4E4')
borde=Border(*[Side(style='thin',color='BFBFBF')]*4)

def enc(ws,t,sub,cols):
    ws['A1']=t; ws['A1'].font=Font(size=14,bold=True,color='1F4E78')
    ws['A2']=sub; ws['A2'].font=Font(size=9,italic=True,color='808080')
    for i,(h,w) in enumerate(cols,1):
        c=ws.cell(row=4,column=i,value=h); c.font=Font(bold=True,color='FFFFFF'); c.fill=azul
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=borde
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[4].height=32; ws.freeze_panes='A5'

def fila(ws,r,v,rell=None,alto=24):
    for i,x in enumerate(v,1):
        c=ws.cell(row=r,column=i,value=x); c.border=borde
        c.alignment=Alignment(vertical='center',wrap_text=True)
        if rell: c.fill=rell
    ws.row_dimensions[r].height=alto

wb=openpyxl.Workbook()
ws=wb.active; ws.title='Panorama'
enc(ws,'Catálogo interno: los 80 publicados y qué le falta a cada uno',
    'Verde = ya liga con su producto real. Ámbar = falta codificarlo, sí entra al almacén. Rojo = no entra al almacén, hay que decidir si debe ser producto.',
    [('Ficha',8),('Clave',13),('Producto',46),('Estado',30),('Qué falta',34),
     ('Existencia',11),('Movs',7),('Flujo',11),('Liga',7)])
r=5
n_ok=n_ent=n_no=0
for f in filas:
    fid, clave, nombre, categ, flujo, ex, movs, liga = f
    if clave:
        estado='Liga con su producto real'; falta='—'; col=verde; n_ok+=1
    elif fid in ENTRAN:
        estado='Sin clave · SÍ entra al almacén'; falta='Codificar como %s' % ENTRAN[fid]; col=ambar; n_ent+=1
    else:
        estado='Sin clave · NO entra al almacén'; falta='Decidir: ¿producto o pedido por descripción? (%s)' % NO_ENTRAN.get(fid,'?'); col=rojo; n_no+=1
    fila(ws,r,[fid, clave or '(sin clave)', nombre, estado, falta, float(ex or 0), int(movs or 0), flujo, liga], col)
    r+=1

ws2=wb.create_sheet('Resumen')
enc(ws2,'El panorama en cuatro renglones','Lo que decide si el catálogo sirve para recibir material.',
    [('Grupo',38),('Cuántos',11),('Qué significa',92)])
RES=[('Ya ligan con su producto real','35','Tienen clave. Al recibir, el material entra al producto correcto. Aquí no falta nada.'),
     ('Sin clave, SÍ entran al almacén','21','Botellas, tubos, puntas, placas, probetas, básculas, termómetros. Necesitan clave o la recepción no se liga a nada. Ya tienen clave propuesta.'),
     ('Sin clave, NO entran al almacén','24','Sillas, escritorios, laptop, impresora, hidrolavadora, reloj. Nadie lleva existencia de una silla. Falta decidir si deben ser producto o pedirse por descripción libre.'),
     ('TOTAL publicados','80','De 1,049 productos del catálogo completo.')]
for n,f in enumerate(RES):
    col = verde if f[1]=='35' else (ambar if f[1]=='21' else (rojo if f[1]=='24' else None))
    fila(ws2,5+n,list(f),col,42)

ws3=wb.create_sheet('Las 2 salidas de los 24')
enc(ws3,'Los 24 que no entran al almacén: las dos salidas','Ninguna es obviamente mejor; dependen de si Compras necesita historial por producto.',
    [('Salida',30),('Qué implica',60),('A favor',52),('En contra',52)])
OPC=[('Codificarlos como producto',
      'Hace falta una clasificación nueva: hoy no existe ni una silla, escritorio o laptop con clave en todo el catálogo.',
      'Quedan en el catálogo con foto, se piden con un clic y generan entrada al almacén. Queda historial de cuántas se han comprado.',
      'Le das de alta producto y existencia a una silla. Es lo que nos metió en este problema: fichas que existen solo para verse.'),
     ('Sacarlos y pedirlos por descripción',
      'El módulo ya lo permite: el renglón de la solicitud acepta texto libre sin producto.',
      'No hace falta clave ni clasificación nueva. No se inventa producto. La solicitud llega igual a Compras.',
      'No genera orden de recepción, porque no hay nada que meter al almacén. Y se pierde el historial por producto.')]
for n,o in enumerate(OPC): fila(ws3,5+n,list(o),ambar,74)

wb.save(B+'Catalogo_panorama_2026-09-29.xlsx')
print('OK: %d ligan, %d por codificar, %d por decidir' % (n_ok, n_ent, n_no))
