# -*- coding: utf-8 -*-
"""Lo que falta producto por producto en el Marketplace Interno."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
B='/tmp/claude-1001/-home-agentia-odoo-odoo-workspace/40e1928e-dd87-4d8e-b8a3-401e8c22dfcd/scratchpad/'

# ficha -> (clave propuesta, grupo)
CLAVE = {
 '2185':('COHIE01','Consumible / embalaje'),'2148':('COESP01','Consumible / limpieza'),
 '2187':('COESC01','Consumible / papelería'),'2176':('COPLA01','Consumible / papelería'),
 '2167':('COTIJ01','Consumible / papelería'),'2164':('COCZP01','Consumible / protección'),
 '2200':('COBOT01','Consumible de laboratorio'),'2201':('COBOT02','Consumible de laboratorio'),
 '2175':('COJER51','Consumible de laboratorio'),'2178':('COMIP01','Consumible de laboratorio'),
 '2153':('COPLP04','Consumible de laboratorio'),'2174':('COPRB06','Consumible de laboratorio'),
 '2165':('COPUN06','Consumible de laboratorio'),'2163':('COTUB09','Consumible de laboratorio'),
 '2157':('EQBAD02','Equipo de laboratorio'),'2207':('EQCEN01','Equipo de laboratorio'),
 '2206':('EQTHI01','Equipo de laboratorio'),'2205':('EQTHI02','Equipo de laboratorio'),
 '2197':('EQTRV02','Equipo de laboratorio'),'2155':('EQVAG01','Equipo de laboratorio'),
 '2159':('MPREC100','Materia prima / medio'),
}
GRUPO_SIN_CASA = {
 '2202':'Mobiliario','2172':'Mobiliario','2203':'Mobiliario','2194':'Mobiliario',
 '2208':'Mobiliario','2198':'Mobiliario','2193':'Mobiliario','2190':'Mobiliario','2184':'Mobiliario',
 '2199':'Cómputo','2196':'Cómputo','2150':'Cómputo','2168':'Cómputo','2156':'Cómputo','2204':'Cómputo',
 '2191':'Instalaciones','2182':'Instalaciones','2322':'Instalaciones','2192':'Instalaciones',
 '2195':'Mantenimiento','2181':'Mantenimiento','2189':'Mantenimiento','2169':'Mantenimiento','2170':'Mantenimiento',
}
PROP_NOTA = {
 'Bata desechable azul':('DECIDE MERY','Existe COBDE01 "Bata Desechable", sin color. ¿El azul importa?'),
 'Etanol':('DECIDE MERY','Existe MPREC18 "Etanol 96%". ¿Sirve el de 96%?'),
 'Biotina':('DECIDE MERY','No se sabe para qué se pide. Hay hoja maestra de Biotina y anticuerpo anti-biotina, no el reactivo'),
 'Glicerol':('DUPLICADA','Es la misma que "Glicerol / Glicerina". Cerrar una de las dos'),
 'Glicerol / Glicerina':('DUPLICADA','Es la misma que "Glicerol". Cerrar una de las dos'),
 'Tinta solvente secado rápido':('ALTA NUEVA','Le falta categoría: sin ella el botón de crear producto no funciona'),
}
azul=PatternFill('solid',fgColor='1F4E78'); ambar=PatternFill('solid',fgColor='FFF2CC')
rojo=PatternFill('solid',fgColor='FCE4E4'); verde=PatternFill('solid',fgColor='E8F4E8')
gris=PatternFill('solid',fgColor='F2F2F2')
borde=Border(*[Side(style='thin',color='BFBFBF')]*4)

def enc(ws,t,sub,cols):
    ws['A1']=t; ws['A1'].font=Font(size=14,bold=True,color='1F4E78')
    ws['A2']=sub; ws['A2'].font=Font(size=9,italic=True,color='808080')
    for i,(h,w) in enumerate(cols,1):
        c=ws.cell(row=4,column=i,value=h); c.font=Font(bold=True,color='FFFFFF'); c.fill=azul
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=borde
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[4].height=30; ws.freeze_panes='A5'

def fila(ws,r,v,rell=None,alto=26):
    for i,x in enumerate(v,1):
        c=ws.cell(row=r,column=i,value=x); c.border=borde
        c.alignment=Alignment(vertical='center',wrap_text=True)
        if rell: c.fill=rell
    ws.row_dimensions[r].height=alto

fichas=[]
for ln in open(B+'f45b.txt'):
    f=ln.rstrip('\n').split('~')
    if len(f)>=5: fichas.append(f)
props=[]
for ln in open(B+'prop15.txt'):
    f=ln.rstrip('\n').split('~')
    if len(f)>=4: props.append(f)

wb=openpyxl.Workbook()
ws=wb.active; ws.title='Empieza aqui'
enc(ws,'Marketplace: lo que falta producto por producto','Al 29-sep-2026, después de fusionar las 15 fichas fantasma y publicar los 15 reactivos.',
    [('Bloque',34),('Cuántos',11),('Qué falta',88)])
RES=[('Fichas con clave propuesta','21','Validar la clave que propuse y que Documentación la libere. 10 necesitan abreviatura nueva.'),
     ('Fichas SIN casa en la normativa','24','Mobiliario, cómputo, instalaciones y mantenimiento. Decidir qué clasificación les toca.'),
     ('Propuestas que son altas de verdad','10','Codificar y dar de alta. Una de ellas no tiene categoría.'),
     ('Propuestas duplicadas entre sí','2','Glicerol y Glicerol / Glicerina son la misma. Cerrar una.'),
     ('Propuestas que esperan tu respuesta','3','Bata AZUL, Etanol de 96% y Biotina.'),
     ('Duplicado muerto en el catálogo','1','STCON01 Control negativo, sin categoría ni existencia, contra STCNL01 con 78 pz.'),
     ('Probetas por confirmar con Luis','2','COPRB05/PTCRS03 y COPRB04/PTCRS04: las de Distribución sin usar.'),
     ('Pares granel/reenvasado','10','MPREC y PTREC con el mismo nombre. No son duplicados, pero no hay cómo distinguirlos al pedir.')]
for n,f in enumerate(RES): fila(ws,5+n,list(f), ambar if f[1] in ('24','3') else None, 30)

ws2=wb.create_sheet('21 con clave propuesta')
enc(ws2,'Fichas que encajan en la codificación actual','Falta tu visto bueno y que Documentación libere la clave.',
    [('Ficha',8),('Producto',46),('Clave propuesta',15),('Grupo',24),('Flujo',12),('Liga',7),('Tu visto bueno',16)])
r=5
for f in fichas:
    if f[0] not in CLAVE: continue
    cl,gr=CLAVE[f[0]]
    fila(ws2,r,[f[0],f[1],cl,gr,f[2],f[3],''], verde); r+=1

ws3=wb.create_sheet('24 sin casa')
enc(ws3,'Fichas que no encajan en ninguna clasificación','No son producto sanitario. Falta decidir qué clasificación les toca.',
    [('Ficha',8),('Producto',48),('Grupo',18),('Flujo',12),('Liga',7),('¿Tiene historial?',15),('Qué clasificación le ponemos',30)])
r=5
for f in fichas:
    if f[0] in CLAVE: continue
    fila(ws3,r,[f[0],f[1],GRUPO_SIN_CASA.get(f[0],'?'),f[2],f[3],f[4],''],
         rojo if f[4]=='SI' else ambar); r+=1

ws4=wb.create_sheet('15 propuestas abiertas')
enc(ws4,'Propuestas que siguen abiertas','Las otras 14 se cerraron el 29-sep: su producto ya existía y ya quedó publicado.',
    [('#',6),('Propuesta',30),('Qué es',14),('Categoría sugerida',26),('Nota',56),('Justificación de quien la pidió',58)])
r=5
for p in props:
    tipo,nota = PROP_NOTA.get(p[1], ('ALTA NUEVA',''))
    col = ambar if tipo=='DECIDE MERY' else (gris if tipo=='DUPLICADA' else verde)
    fila(ws4,r,[p[0],p[1],tipo,p[2],nota,p[3]],col,34); r+=1

ws5=wb.create_sheet('Nombres repetidos')
enc(ws5,'Productos con el mismo nombre en el catálogo','Uno es duplicado real; los demás son pares granel/reenvasado o de Distribución.',
    [('Nombre',26),('Clave A',12),('Qué es A',30),('Clave B',12),('Qué es B',30),('Veredicto',18),('Qué hacer',44)])
DUP=[('Control negativo','STCON01','sin categoría · 0 existencia · 3 movs','STCNL01','78 pz · 18 movs','DUPLICADO REAL','Archivar STCON01, no borrar: tiene 3 movimientos'),
 ('Probeta 1 L','COPRB05','en catálogo · 1 mov','PTCRS03','Distribución · 0 · sin movimientos','CONFIRMAR CON LUIS','¿El de Distribución se usa?'),
 ('Probeta 250 ml','COPRB04','en catálogo · 1 pz','PTCRS04','Distribución · 0 · sin movimientos','CONFIRMAR CON LUIS','¿El de Distribución se usa?'),
 ('CHAPS','MPREC23','granel','PTREC09','reenvasado','No es duplicado','Distinguir en el nombre: (granel) y (reenvasado)'),
 ('DTT','MPREC26','granel · 50 · en 2 recetas','PTREC13','reenvasado · 0','No es duplicado','Distinguir en el nombre'),
 ('Guanidine hydrochloride','MPREC24','granel','PTREC10','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('HEPES','MPREC20','granel','PTREC06','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('IPTG','MPREC25','granel · 1000','PTREC12','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('MES','MPREC22','granel · 2000','PTREC08','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('MOPS','MPREC21','granel','PTREC07','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('TCEP-HCl','MPREC27','granel · 0 · en 3 recetas','PTREC11','reenvasado · nunca vendido','No es duplicado','Distinguir en el nombre'),
 ('Tris base','MPREC10','granel','PTREC05','reenvasado','No es duplicado','Distinguir en el nombre'),
 ('Tris-HCl','MPREC04','granel · 7999','PTREC04','reenvasado','No es duplicado','Distinguir en el nombre')]
for n,d in enumerate(DUP):
    fila(ws5,5+n,list(d), rojo if 'REAL' in d[5] else (ambar if 'LUIS' in d[5] else None), 28)

wb.save(B+'Marketplace_pendiente_2026-09-29.xlsx')
print('OK: %d fichas, %d propuestas, %d nombres repetidos' % (len(fichas), len(props), len(DUP)))
