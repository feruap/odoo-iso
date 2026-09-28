# -*- coding: utf-8 -*-
"""Hoja de trabajo de las soluciones de captura, para que Mery la complete."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# clave, nombre solucion, clave MP, nombre MP, clave proveedor
FILAS = [
    ('SPSCA01', 'Solución de captura anti-influenza A', 'MPANT22', 'Anticuerpo de captura anti-influenza A', 'BRCINFS104'),
    ('SPSCA02', 'Solución de captura anti-influenza B', 'MPANT23', 'Anticuerpo de captura anti-influenza B', 'BRNINFC201'),
    ('SPSCA03', 'Solución de captura anti-THC', '', '', ''),
    ('SPSCA04', 'Solución de captura anti-COC', '', '', ''),
    ('SPSCA05', 'Solución de captura anti-AMP', '', '', ''),
    ('SPSCA06', 'Solución de captura anti-MET', '', '', ''),
    ('SPSCA07', 'Solución de captura anti-OPI', '', '', ''),
    ('SPSCA08', 'Solución de captura anti-biotina', 'MPANT27', 'Anticuerpo de captura anti-biotina', 'GENAMHN101'),
    ('SPSCA09', 'Solución de captura anti-mioglobina', 'MPANT24', 'Anticuerpo de captura anti-myo', ''),
    ('SPSCA10', 'Solución de captura anti-CK-MB', 'MPANT25', 'Anticuerpo de captura anti-ckmb', ''),
    ('SPSCA11', 'Solución de captura anti-cTnI', 'MPANT26', 'Anticuerpo de captura anti-cTnl', ''),
    ('SPSCA12', 'Solución de captura anti-Chlamydia trachomatis', 'MPANT14', 'Anticuerpo de captura anti-clamidia', ''),
    ('SPSCA13', 'Solución de captura IgG anti-humano', 'MPANT07', 'Anticuerpo de captura anti-IgG humano', 'BRJIGGI102'),
    ('SPSCA14', 'Solución de captura IgM anti-humano', 'MPANT08', 'Anticuerpo de captura anti-IgM humano', 'GRCIGMC101'),
    ('SPSCA15', 'Solución de captura anti-proteína N', 'MPANT04', 'Anticuerpo de captura anti-proteína N SARS-CoV-2', 'FPZ0638'),
    ('SPSCA16', 'Solución de captura anti-NS1', 'MPANT28', 'Anticuerpo de captura anti-NS1 del Dengue', ''),
    ('SPSCA17', 'Solución de captura anti-ALC', '', '', ''),
    ('SPSCA18', 'Solución de captura anti-ferritina', 'MPANT42', 'Anticuerpo de captura anti-ferritina', ''),
    ('SPSCA19', 'Solución de captura anti-hCG', 'MPANT09', 'Anticuerpo de captura anti-hCG', 'FPZ0612'),
    ('SPSCA20', 'Solución de captura anti-hemoglobina glicada', 'MPANT12', 'Anticuerpo de captura anti-HbA1c', 'BRJHBAS101'),
    ('SPSCA21', 'Solución de captura anti-PSA', 'MPANT37', 'Anticuerpo de captura anti-PSA', 'BENPSAN101'),
    ('SPSCA22', 'Solución de captura anti-rotavirus', 'MPANT44', 'Anticuerpo de captura anti-Rotavirus', ''),
    ('SPSCA23', 'Solución de captura anti-adenovirus', 'MPANT46', 'Anticuerpo de captura anti-Adenovirus', ''),
    ('SPSCA24', 'Solución de captura anti-RSV', 'MPANT39', 'Anticuerpo de captura anti-RSV', 'FPZ0181'),
    ('SPSPC01', 'Solución de captura Antígeno de Treponema pallidum', 'MPAG09', 'Antígeno de captura Treponema pallidum', 'GRCTPS206'),
    ('SPSAN01', 'Solución de captura 25 (OH) D', 'MPAG18', 'Antígeno Vitamina D-BSA', 'FPZ0188-3'),
    ('SPSCA25', 'Solución de captura anti-ratón', 'MPANT03', 'Anticuerpo policlonal anti-ratón (línea control)', 'GRCGAMS002'),
    ('SPSCA26', 'Solución de captura anti-p24 del VIH 1', 'MPANT41', 'Anticuerpo de captura anti-p24', ''),
    ('SPSCA27', 'Solución de captura anti-TSH', 'MPANT55', 'Anticuerpo de captura anti-TSH', 'BECTSHS102'),
    ('SPSCA28', 'Solución de captura anti-Salmonella typhi', 'MPANT49', 'Anticuerpo de captura anti-Salmonella', ''),
]

# Las de drogas y la de alcohol van sin MP a proposito: Mery indico el 28-sep
# que por ahora no se fabrican desde este punto, solo por linea corta.
SIN_MP = {'SPSCA03', 'SPSCA04', 'SPSCA05', 'SPSCA06', 'SPSCA07', 'SPSCA17'}

ENCABEZADOS = [
    ('Solución', 46), ('Clave', 11),
    ('Anticuerpo / antígeno', 46), ('Clave Amunet', 14), ('Clave proveedor', 16),
    ('Cantidad de anticuerpo', 22), ('Solución diluyente', 26),
    ('Impresión en membrana', 22),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Soluciones de captura'

azul = PatternFill('solid', fgColor='1F4E78')
ambar = PatternFill('solid', fgColor='FFF2CC')
gris = PatternFill('solid', fgColor='F2F2F2')
borde = Border(*[Side(style='thin', color='BFBFBF')] * 4)

ws['A1'] = 'Soluciones de captura — línea de laminado'
ws['A1'].font = Font(size=14, bold=True, color='1F4E78')
ws['A2'] = ('Se imprimen en la membrana. Caducidad 3 días, no piden análisis. '
            'Las columnas ámbar son las que faltan por definir.')
ws['A2'].font = Font(size=9, italic=True, color='808080')

FILA0 = 4
for i, (titulo, ancho) in enumerate(ENCABEZADOS, start=1):
    c = ws.cell(row=FILA0, column=i, value=titulo)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = azul
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = borde
    ws.column_dimensions[get_column_letter(i)].width = ancho
ws.row_dimensions[FILA0].height = 32

for n, (clave, sol, mp, nombre_mp, prov) in enumerate(FILAS):
    r = FILA0 + 1 + n
    valores = [sol, clave, nombre_mp, mp, prov, None, None, None]
    for i, v in enumerate(valores, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = borde
        c.alignment = Alignment(vertical='center', wrap_text=(i in (1, 3)))
        if i >= 6:
            c.fill = ambar
        elif clave in SIN_MP and i in (3, 4, 5):
            c.fill = gris
    ws.row_dimensions[r].height = 26

r = FILA0 + len(FILAS) + 2
ws.cell(row=r, column=1, value='Nota').font = Font(bold=True)
ws.cell(row=r + 1, column=1,
        value='Las 5 de drogas (THC, COC, AMP, MET, OPI) y la de alcohol van sin anticuerpo: '
              'por ahora no se fabrican desde este punto, solo por línea corta.')
ws.cell(row=r + 2, column=1,
        value='SPSPC01 comparte abreviatura con las soluciones de borato (SPSPC03/04/05); '
              'Documentación tiene que confirmar si le toca clave propia.')
ws.cell(row=r + 3, column=1,
        value='SPSCA25 anti-ratón es la LÍNEA DE CONTROL: la llevan las 15 hojas que se '
              'laminan aquí, no una sola. Su cantidad y volumen aplican a todas.')
ws.cell(row=r + 4, column=1,
        value='Falta la línea de referencia de las dos semicuantitativas (SPHMC38 PSA y '
              'SPHMC52 TSH): qué anticuerpo y a qué concentración.')
for k in (1, 2, 3, 4):
    ws.cell(row=r + k, column=1).font = Font(size=9, italic=True)

ws.freeze_panes = 'A%d' % (FILA0 + 1)
wb.save('/tmp/claude-1001/-home-agentia-odoo-odoo-workspace/40e1928e-dd87-4d8e-b8a3-401e8c22dfcd/scratchpad/Soluciones_de_captura.xlsx')
print('OK %d filas' % len(FILAS))
